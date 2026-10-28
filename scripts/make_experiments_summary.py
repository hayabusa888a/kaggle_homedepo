"""実験の振り返り用の Excel（experiments_summary.xlsx）を作る。実験を追加したら rows に 1 行足して実行し直す。"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "experiments_summary.xlsx"
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
BEST_FILL = PatternFill("solid", fgColor="E2EFDA")
NG_FILL = PatternFill("solid", fgColor="FCE4D6")
REF_FILL = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")


def f(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = f(True, "FFFFFF"), HEAD_FILL, Alignment(wrap_text=True, vertical="center"), BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


wb = Workbook()

# ---------------------------------------------------------------- 実験一覧
ws = wb.active
ws.title = "実験一覧"
cols = ["No", "実験", "日付", "カテゴリ", "内容", "比較対象", "CV RMSE", "Private LB", "Public LB",
        "LB−CV\n(Private)", "CV の変化\n(比較対象から)", "Private LB の変化\n(比較対象から)", "判定", "メモ・学び", "スクリプト / 出力"]
widths = [5, 20, 11, 15, 52, 20, 10, 10, 10, 10, 12, 13, 12, 60, 38]
header(ws, 1, cols, widths)

N = None
rows = [
    # 実験, 日付, カテゴリ, 内容, 比較対象, CV, Private, Public, 判定, メモ, スクリプト
    ("exp001", "2026-10-13", "ベースライン", "TF-IDF コサイン類似度（検索語×タイトル、検索語×説明文）＋ LightGBM", N, 0.50579, 0.50782, N, "採用（基準）", "特徴量 2 つだけ。CV の分け方は検索語単位の GroupKFold（test の 73% は train にない検索語）", "exp/exp001_tfidf_lgbm.py"),
    ("exp002", "2026-10-15", "特徴量", "＋ 検索語・タイトルの TF-IDF ベクトル（疎行列のまま）", "exp001", 0.49581, 0.49333, N, "採用", "for/in/to など「検索語の形」や商品カテゴリの単語が効いた", "exp/exp002_tfidf_vec_lgbm.py"),
    ("exp003", "2026-10-15", "前処理", "＋ 表記の正規化（連結語・単位・寸法・ハイフン）＋ Snowball ステミング", "exp002", 0.48587, N, N, "採用", "語彙が 1/3 に。1 文字の数字を残したことも効いた", "exp/exp003_preprocess.py"),
    ("exp004", "2026-10-15", "スペル修正", "＋ 検索語のスペル修正（SymSpell、商品テキストの語彙を辞書に）", "exp003", 0.48211, 0.47993, N, "採用", "修正が入るのは 14% の行。ブランド名に強い", "exp/exp004_spell_symspell.py"),
    ("exp005", "2026-10-15", "スペル修正", "exp003 ＋ 検索語のスペル修正（ローカル LLM：Llama 3.1 8B）", "exp004", 0.48250, 0.48024, N, "不採用", "SymSpell とほぼ同等。修正結果の一致は 55% で得意分野が違う（LLM は一般語・文脈、SymSpell は商品の語彙）", "exp/exp005_spell_llm.py"),
    ("exp006", "2026-10-16", "特徴量", "＋ 単語一致の数え上げ（18 個）＋ ブランド一致（4 個）", "exp004", 0.46648, 0.46501, N, "採用", "単純な「何割一致したか」が 1 位に。最後の単語の一致、別ブランドの言及が効いた", "exp/exp006_match_brand.py"),
    ("exp007", "2026-10-16", "特徴量", "＋ BM25 ＋ 文字 n-gram の類似度 ＋ 属性値との一致", "exp006", 0.46387, 0.46273, N, "採用", "文字 n-gram が 2 位。属性値はほぼ効かない", "exp/exp007_bm25_char_attr.py"),
    ("exp008", "2026-10-16", "特徴量", "＋ ブランド名の TF-IDF（類似度 ＋ ベクトル）", "exp007", 0.46364, N, N, "不採用", "誤差の範囲。exp006 のブランド特徴量・タイトルと情報が重なる", "exp/exp008_brand_tfidf.py"),
    ("exp009", "2026-10-16", "埋め込み", "exp007 ＋ 埋め込み（nomic-embed-text、Ollama）の類似度 ＋ ベクトル", "exp007", 0.45262, 0.45013, N, "採用", "意味の近さ（同義語）を拾える。単独の相関 0.428 は当時の最高", "exp/exp009_embedding.py"),
    ("exp010", "2026-10-16", "埋め込み", "exp007 ＋ 埋め込み（mxbai-embed-large、GPU で直接計算）＋ 説明文のベクトルも", "exp009", 0.44584, 0.44345, N, "採用", "PyTorch で GPU を直接使い 84 分 → 33 分", "exp/exp010_embedding_mxbai.py"),
    ("exp011", "2026-10-17", "検証", "exp009（nomic）＋ 説明文の埋め込みベクトル", "exp009", 0.45252, N, N, "検証", "説明文のベクトルは効果なし → exp010 の改善はモデルの変更によるもの。gain が大きくても予測を良くしているとは限らない", "exp/exp011_nomic_desc_vec.py"),
    ("exp012", "2026-10-17", "検証", "exp010 から説明文の埋め込みベクトルを外す", "exp010", 0.44593, N, N, "採用", "精度は同じで列数・時間を削減", "exp/exp012_mxbai_no_desc_vec.py"),
    ("exp013", "2026-10-17", "リランカー", "＋ cross-encoder（ms-marco-MiniLM-L-6-v2）のスコア", "exp012", 0.44496, N, N, "不採用", "軽量モデルでは効果小。exp014 に置き換え", "exp/exp013_rerank_minilm.py"),
    ("exp014", "2026-10-17", "リランカー", "＋ cross-encoder（mxbai-rerank-base-v1、DeBERTa）のスコア（学習前のまま）", "exp012", 0.44025, N, N, "採用", "説明文のスコアで MiniLM に大差", "exp/exp014_rerank_mxbai.py"),
    ("exp015", "2026-10-17", "リランカー", "exp014 ＋ MiniLM のスコアも両方", "exp014", 0.44011, N, N, "不採用", "同じ種類の情報で上乗せなし", "exp/exp015_rerank_both.py"),
    ("exp016", "2026-10-17", "チューニング", "LightGBM のハイパーパラメータを Optuna で調整（30 回）＋ learning_rate 0.02", "exp014", 0.43841, 0.43710, N, "採用", "小さく慎重な木 + 強い L1 + feature_fraction 0.31。改善の大半はパラメータ（lr の効果は −0.0003）", "exp/exp016_tuning.py"),
    ("exp017", "2026-10-18", "特徴量", "＋ 同じ検索語の中で比べた相対的な特徴量（平均・最大との差、順位）", "exp016", 0.43582, 0.43452, N, "採用", "手軽で効果大。q_size（検索語の商品数）はデータの作り方由来の偏り", "exp/exp017_query_relative.py"),
    ("exp018", "2026-10-18", "ファインチューニング", "＋ ファインチューニングした cross-encoder（bge-reranker-base、検索語×タイトル、out-of-fold）", "exp017", 0.42836, 0.42810, N, "採用", "単独で相関 0.554。DeBERTa（mxbai-rerank）は transformers 5.x で NaN になり断念", "exp/exp018_finetune_ce.py"),
    ("exp019", "2026-10-18", "特徴量", "＋ 検索語とタイトルの主語（商品の種類）の一致", "exp018", 0.42835, N, N, "不採用", "誤差分析の「付属品の取り違え」は、実はラベルのノイズだった", "exp/exp019_head_match.py"),
    ("exp020", "2026-10-18", "ファインチューニング", "＋ ファインチューニング（bge-reranker-base、検索語×タイトル + 説明文）", "exp018", 0.42615, 0.42668, N, "採用", "説明文を入れると単独で −0.0046。この頃から LB が CV より悪くなり始める", "exp/exp020_finetune_ce_desc.py"),
    ("exp021", "2026-10-19", "ファインチューニング", "＋ 大きいモデル（bge-reranker-v2-m3、タイトル）", "exp020", 0.42283, N, N, "採用", "単語の埋め込み層の固定 + 8bit AdamW で 8GB に収めた（約 8 倍速）", "exp/exp021_finetune_ce_large.py"),
    ("exp022", "2026-10-19", "ファインチューニング", "＋ 大きいモデル（v2-m3、タイトル + 説明文）", "exp021", 0.42098, 0.42315, N, "採用", "予測を fp32 でまとめて行うとメモリがあふれ停滞 → bf16 予測に修正", "exp/exp022_finetune_ce_large_desc.py"),
    ("exp023", "2026-10-19", "アンサンブル", "24 メンバー（LGB/XGB/Cat/Ridge/RF/ET/kNN/MLP）を OOF で重み最適化", "exp022", 0.41938, 0.42338, N, "不採用", "CV では改善したが LB は悪化。重みは Ridge・MLP・RF など LGB と性質の違うモデルに集中", "exp/exp023_blend.py"),
    ("exp024_lgb", "2026-10-19", "検証・対策", "test のファインチューニングのスコアを fold ごとのモデルの予測にそろえる（LightGBM 単体）", "exp022", 0.42099, 0.42257, 0.42444, "採用", "train は 1 モデルの予測、test は 5 モデルの平均というずれを解消。LB −0.0006", "output/exp024/submission_lgb.csv"),
    ("exp024_blend", "2026-10-19", "検証・対策", "同上（アンサンブル）", "exp023", 0.41920, 0.42282, 0.42485, "不採用", "対策で LB −0.0006 だが LGB 単体に及ばず", "exp/exp024_blend.py"),
    ("exp025_lgb", "2026-10-20", "ファインチューニング", "v2-m3（タイトル + 説明文）を 2 つの種で平均（LightGBM 単体）", "exp024_lgb", 0.42030, N, N, "採用", "単独 0.43708 → 0.43407。8 時間の不在中に独立プロセスで実行", "output/exp025/submission_lgb.csv"),
    ("exp026_lgb", "2026-10-20", "ファインチューニング", "v2-m3（タイトルのみ）も 2 つの種で平均（LightGBM 単体）", "exp025_lgb", 0.42015, 0.42254, 0.42439, "採用", "種の平均は CV では効くが LB にはほぼ効かず", "output/exp026/submission_lgb.csv"),
    ("exp026_blend", "2026-10-20", "アンサンブル", "同上（アンサンブル）", "exp024_blend", 0.41854, 0.42268, 0.42461, "不採用", "LGB 単体に及ばず", "exp/exp026_run.py"),
    ("exp027", "2026-10-20", "特徴量選択", "LightGBM で数値の特徴量 80 列だけ使う（埋め込み・TF-IDF ベクトルを外す）", "exp026_lgb", 0.41776, 0.42354, 0.42547, "不採用", "CV −0.0024 なのに LB +0.0010。ファインチューニングのスコアへの依存を強めると LB で悪化", "output/exp027/submission_num.csv"),
    ("exp028_all", "2026-10-21", "検証・対策", "train にある検索語の test 行は、その検索語を見ていない fold のモデルのスコアを使う", "exp026_lgb", 0.42015, 0.42340, 0.42529, "不採用", "仮説は外れ。見た検索語の高めのスコアは正しい情報だった", "output/exp028/submission_all.csv"),
    ("exp028_num", "2026-10-21", "検証・対策", "同上（数値の特徴量だけ）", "exp027", 0.41776, 0.42434, 0.42636, "不採用", "同上", "output/exp028/submission_num.csv"),
    ("exp029_all", "2026-10-27", "ファインチューニング", "＋ ファインチューニングした DeBERTa（mxbai-rerank-base-v1、タイトル + 説明文、transformers 4.49）", "exp026_lgb", 0.41812, 0.42011, 0.42207, "採用", "1.8 億で v2-m3（5.7 億）と同精度。系統の違うモデルは LB でもしっかり効く。学習は 7 時間（GPU メモリがあふれ気味）", "output/exp029/submission_all.csv"),
    ("exp029_num", "2026-10-27", "特徴量選択", "同上（数値の特徴量 84 列だけ、参考）", "exp029_all", 0.41547, N, N, "参考", "未提出", "output/exp029/submission_num.csv"),
    ("exp030_all", "2026-10-28", "ファインチューニング", "＋ ファインチューニングした DeBERTa（検索語 × タイトル）", "exp029_all", 0.41776, 0.41998, 0.42198, "採用", "LB はわずかに改善（Public −0.00009）。単独 0.43951 はタイトル版で最良。ただし DeBERTa（タイトル + 説明文）との相関 0.945 で、CV の上乗せは −0.0004 と小さい。学習は 1 時間", "output/exp030/submission_all.csv"),
    ("exp030_num", "2026-10-28", "特徴量選択", "同上（数値の特徴量 88 列だけ、参考）", "exp030_all", 0.41521, N, N, "参考", "未提出", "output/exp030/submission_num.csv"),
    ("exp031_all", "2026-10-28", "ファインチューニング", "exp029 ＋ ファインチューニングした ModernBERT（gte-reranker-modernbert-base、タイトル + 説明文）", "exp029_all", 0.41743, 0.41990, 0.42184, "採用（ベスト）", "LB で Public −0.00023、新しいベスト（exp030 とは別に exp029 に加えたもの）。別系統のリランカー。単独 0.43602 は全モデルで最良（1.5 億パラメータ）。ほかとの相関 0.87〜0.93。学習は 2 時間 40 分", "output/exp031/submission_all.csv"),
    ("exp031_num", "2026-10-28", "特徴量選択", "同上（数値の特徴量 88 列だけ、参考）", "exp031_all", 0.41496, N, N, "参考", "未提出", "output/exp031/submission_num.csv"),
]
first = 2
for i, r in enumerate(rows):
    rr = first + i
    exp, date, cat, desc, base, cv, priv, pub, judge, memo, path = r
    vals = [i + 1, exp, date, cat, desc, base, cv, priv, pub]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=rr, column=j, value=v)
    # 数式: LB−CV、比較対象からの変化（比較対象の行を名前で探す）
    ws.cell(row=rr, column=10, value=f'=IF(H{rr}="","",H{rr}-G{rr})')
    ws.cell(row=rr, column=11, value=f'=IF(F{rr}="","",G{rr}-INDEX($G:$G,MATCH(F{rr},$B:$B,0)))')
    ws.cell(row=rr, column=12, value=f'=IF(OR(F{rr}="",H{rr}=""),"",IFERROR(IF(INDEX($H:$H,MATCH(F{rr},$B:$B,0))="","",H{rr}-INDEX($H:$H,MATCH(F{rr},$B:$B,0))),""))')
    ws.cell(row=rr, column=13, value=judge)
    ws.cell(row=rr, column=14, value=memo)
    ws.cell(row=rr, column=15, value=path)
    fill = BEST_FILL if "ベスト" in judge else NG_FILL if judge == "不採用" else REF_FILL if judge in ("検証", "参考") else None
    for j in range(1, 16):
        c = ws.cell(row=rr, column=j)
        c.font, c.border, c.alignment = f(), BORDER, WRAP
        if fill:
            c.fill = fill
        if j in (7, 8, 9):
            c.number_format = "0.00000"
            if j in (8, 9) and c.value is not None:
                c.font = f(color="0000FF")
        if j in (10, 11, 12):
            c.number_format = "+0.0000;-0.0000;0.0000"
last_row = first + len(rows) - 1
ws.cell(row=last_row + 2, column=2, value="凡例").font = f(True)
notes = [
    ("緑の行", "現在のベスト（exp031_all）"), ("オレンジの行", "不採用（以降の実験の土台にしていない）"), ("灰色の行", "検証・参考のための実験"),
    ("青い数字", "Kaggle に提出して得た LB（ユーザーから共有された値）。空欄は未提出、または Public LB を記録していない"),
    ("LB−CV・変化の列", "数式で計算（比較対象は「比較対象」列の実験名を「実験」列から探す）。LB の変化は、両方に Private LB があるときだけ表示"),
    ("CV", "検索語単位の GroupKFold（5 分割）の RMSE。exp023〜は入れ子の CV（アンサンブル）または LightGBM 単体の OOF"),
]
for k, (a, b) in enumerate(notes):
    ws.cell(row=last_row + 3 + k, column=2, value=a).font = f()
    ws.cell(row=last_row + 3 + k, column=5, value=b).font = f()
ws.auto_filter.ref = f"A1:O{last_row}"

# ---------------------------------------------------------------- 概要
ov = wb.create_sheet("概要", 0)
ov.column_dimensions["A"].width = 26
ov.column_dimensions["B"].width = 95
items = [
    ("Home Depot Product Search Relevance：実験の振り返り（exp001〜exp031）", None, "title"),
    ("タスク", "検索語と商品のペアの関連度 relevance（1〜3、複数人の評価の平均）を予測する。評価指標は RMSE", None),
    ("期間", "2026-10-13 〜 2026-10-28", None),
    ("現在のベスト", "exp031_all：LightGBM 単体（全部の列）。CV 0.41743 / Private LB 0.41990 / Public LB 0.42184", "best"),
    ("最初のベースライン", "exp001：CV 0.50579 / Private LB 0.50782（ベストまで −0.0879）", None),
    ("ベストの構成", "前処理（正規化・SymSpell・ステミング）→ TF-IDF・BM25・単語一致・文字 n-gram・ブランド・属性値 → 埋め込み（mxbai-embed-large）→ 学習前の cross-encoder（mxbai-rerank-base）→ ファインチューニングした cross-encoder 6 種（bge-reranker-base ×2、v2-m3 ×2（種の平均）、DeBERTa ×1、ModernBERT ×1）→ 同じ検索語の中での相対値 → LightGBM（Optuna で調整、lr 0.02）", None),
    ("CV の方針", "検索語単位の GroupKFold（5 分割）。test の 73% の行は train にない検索語で、train では 1 つの検索語が平均 6〜7 行出てくるため", None),
    ("データの要点", "test の 57% の行は train にない商品。11.6% は train のどの商品より新しい（product_uid > 206,650）。新しい商品ほど relevance が低く、誤差が大きい。同じ検索語でほぼ同じ商品の組でも評価のばらつきが RMSE 換算で 0.288 あり、ラベルのノイズが大きい", None),
    ("最大の課題", "ファインチューニングのスコアへの依存を強める変更（アンサンブル、種の平均、数値だけの特徴量）ほど、CV は良くなるのに LB は良くならない。CV では train より新しい商品を評価できない。採否は Public LB で判断している", None),
    ("環境", "RTX 5070 Laptop（8GB、電源アダプタ次第で 25W 制限）。.venv-gpu（transformers 5.20）と .venv-deberta（transformers 4.49、DeBERTa の学習用）。長時間の処理はツールの 2 時間制限のため独立プロセスで実行", None),
    ("詳細", "各実験の詳細は EXPERIMENTS.md、コードは exp/・src/・scripts/", None),
    ("シートの説明", None, "head"),
    ("実験一覧", "全実験の CV・LB・比較対象からの変化（数式）・判定・メモ", None),
    ("ファインチューニング比較", "ファインチューニングした cross-encoder ごとの単独の精度と時間", None),
    ("学び", "実験から得た教訓と、その根拠になった実験", None),
    ("CV と LB", "提出した実験の CV と Private LB の推移（表とグラフ）", None),
]
r = 1
for a, b, kind in items:
    ca = ov.cell(row=r, column=1, value=a)
    if kind == "title":
        ca.font = f(True, size=14)
        r += 2
        continue
    if kind == "head":
        r += 1
        ov.cell(row=r, column=1, value=a).font = f(True, size=12)
        ov.cell(row=r - 1, column=1).value = None
        r += 1
        continue
    ca.font, ca.alignment = f(True), WRAP
    cb = ov.cell(row=r, column=2, value=b)
    cb.font, cb.alignment = f(), WRAP
    if kind == "best":
        ca.fill = cb.fill = BEST_FILL
    r += 1

# ---------------------------------------------------------------- ファインチューニング比較
ft = wb.create_sheet("ファインチューニング比較")
hdr = ["モデル", "系統", "パラメータ数", "入力", "学習の工夫", "単独の OOF RMSE", "relevance との相関", "学習時間（5 fold）", "使用した実験", "備考"]
header(ft, 1, hdr, [26, 14, 12, 18, 34, 14, 14, 16, 16, 46])
ftrows = [
    ("bge-reranker-base", "XLM-RoBERTa", "2.8 億", "タイトル", "—", 0.44997, 0.554, "29 分", "exp018〜", ""),
    ("bge-reranker-base", "XLM-RoBERTa", "2.8 億", "タイトル + 説明文（256）", "—", 0.44534, 0.567, "178 分", "exp020〜", ""),
    ("bge-reranker-v2-m3", "XLM-RoBERTa", "5.7 億", "タイトル", "埋め込み層を固定 + 8bit AdamW", 0.44113, 0.574, "65 分", "exp021〜", ""),
    ("bge-reranker-v2-m3（2 つの種の平均）", "XLM-RoBERTa", "5.7 億 ×2", "タイトル", "同上", 0.43792, N, "+66 分", "exp026〜", "種の相関 0.958"),
    ("bge-reranker-v2-m3", "XLM-RoBERTa", "5.7 億", "タイトル + 説明文（256）", "同上 + バッチ 8 × 4 + bf16 予測", 0.43708, 0.585, "283 分", "exp022〜", ""),
    ("bge-reranker-v2-m3（2 つの種の平均）", "XLM-RoBERTa", "5.7 億 ×2", "タイトル + 説明文（256）", "同上", 0.43407, N, "+4.7 時間", "exp025〜", "種の相関 0.964"),
    ("mxbai-rerank-base-v1", "DeBERTa-v3", "1.8 億", "タイトル（128）", "transformers 4.49 の別環境。バッチ 16 × 2", 0.43951, N, "1 時間", "exp030", "DeBERTa（タイトル + 説明文）との相関 0.945"),
    ("mxbai-rerank-base-v1", "DeBERTa-v3", "1.8 億", "タイトル + 説明文（256）", "transformers 4.49 の別環境。バッチ 16 × 2", 0.43702, N, "7 時間", "exp029", "transformers 5.x では NaN。ほかのスコアとの相関 0.88〜0.92"),
    ("gte-reranker-modernbert-base", "ModernBERT", "1.5 億", "タイトル + 説明文（256）", "—（.venv-gpu のまま）", 0.43602, 0.582, "2 時間 40 分", "exp031", "ほかのスコアとの相関 0.87〜0.93"),
    ("（学習前）mxbai-rerank-base-v1", "DeBERTa-v3", "1.8 億", "タイトル", "ファインチューニングなし", N, 0.433, "51 分（採点のみ）", "exp014〜", "参考"),
]
for i, rowv in enumerate(ftrows, 2):
    for j, v in enumerate(rowv, 1):
        c = ft.cell(row=i, column=j, value=v)
        c.font, c.border, c.alignment = f(), BORDER, WRAP
        if j == 6:
            c.number_format = "0.00000"
        if j == 7:
            c.number_format = "0.000"

# ---------------------------------------------------------------- 学び
lw = wb.create_sheet("学び")
header(lw, 1, ["No", "学び", "根拠の実験"], [5, 100, 30])
lessons = [
    ("CV の分け方は test との関係に合わせる。test の 73% が未知の検索語なので、検索語単位の GroupKFold にした", "exp001〜"),
    ("単純な特徴量（単語一致の割合、最後の単語の一致など）が、複雑な類似度より効くことがある", "exp006"),
    ("意味の近さ（埋め込み）と、ペアを同時に読む cross-encoder は大きく効く。cross-encoder はファインチューニングでさらに伸びる", "exp009, exp014, exp018"),
    ("gain（重要度）が大きい特徴量でも、予測を良くしているとは限らない。入れた場合と外した場合を比べて確かめる", "exp011, exp012"),
    ("誤差分析で目立った「付属品の取り違え」の多くは、実はラベルのノイズだった。大きく外した行は、同じ商品群の評価のばらつきも確かめる", "exp019"),
    ("ラベルを使う特徴量（スタッキング）は、2 段目と同じ fold で out-of-fold にしないとリークする。fold の一致はコードで確かめる", "exp018"),
    ("train と test で特徴量の作り方をそろえる（OOF は 1 モデル、test も fold ごとに 1 モデル）と LB が改善した", "exp024"),
    ("ファインチューニングのスコアへの依存を強める変更（アンサンブル、種の平均、数値だけ）は、CV では良くても LB では効かない・悪化する。CV では train より新しい商品を評価できない", "exp023, exp025〜exp028"),
    ("一方、系統の違うモデル（DeBERTa）を加えると、CV 以上に LB が改善した。同じ系統を重ねるより、系統を増やすほうが頑健", "exp029"),
    ("系統の違うモデルを加える効果は、加えるほど小さくなる（DeBERTa は LB −0.0024、次の ModernBERT は −0.0002）。同じ系統で入力だけ変えた版（DeBERTa タイトル）はさらに小さい", "exp029, exp030, exp031"),
    ("仮説（見た検索語のスコアの偏り）は LB で否定された。仮説は小さく試し、外れたら記録して戻す", "exp028"),
    ("Private LB を見て何度も選ぶと、test への過学習になる。採否は Public LB で判断する", "exp027〜"),
    ("8GB の GPU では、単語の埋め込み層の固定・8bit AdamW・バッチの分割・bf16 予測でメモリを抑えると、大きいモデルも現実的な時間で学習できる。あふれると数倍遅くなる", "exp021, exp022, exp029"),
    ("ライブラリのバージョンでモデルが学習できないことがある（DeBERTa は transformers 5.x で NaN、4.49 では学習できた）", "exp018, exp029"),
    ("長時間の処理は、ツールの 2 時間制限を避けて独立プロセスで実行し、fold ごとに保存して途中から再開できるようにする", "exp025, exp026, exp029"),
]
for i, (a, b) in enumerate(lessons, 2):
    for j, v in enumerate([i - 1, a, b], 1):
        c = lw.cell(row=i, column=j, value=v)
        c.font, c.border, c.alignment = f(), BORDER, WRAP

# ---------------------------------------------------------------- CV と LB
cl = wb.create_sheet("CV と LB")
header(cl, 1, ["実験", "CV RMSE", "Private LB", "LB−CV"], [16, 12, 12, 12])
cl.freeze_panes = "A2"
submitted = [i for i, r in enumerate(rows) if r[6] is not None]
for k, i in enumerate(submitted, 2):
    src = first + i
    cl.cell(row=k, column=1, value=f"=実験一覧!B{src}")
    cl.cell(row=k, column=2, value=f"=実験一覧!G{src}")
    cl.cell(row=k, column=3, value=f"=実験一覧!H{src}")
    cl.cell(row=k, column=4, value=f"=C{k}-B{k}")
    for j in range(1, 5):
        c = cl.cell(row=k, column=j)
        c.font, c.border = f(color="008000" if j < 4 else "000000"), BORDER
        if j in (2, 3):
            c.number_format = "0.00000"
        if j == 4:
            c.number_format = "+0.0000;-0.0000;0.0000"
n_sub = len(submitted)
chart = LineChart()
chart.title = "提出した実験の CV と Private LB"
chart.y_axis.title = "RMSE"
chart.y_axis.number_format = "0.000"
chart.height, chart.width = 10, 24
data = Reference(cl, min_col=2, max_col=3, min_row=1, max_row=1 + n_sub)
chart.add_data(data, titles_from_data=True)
chart.set_categories(Reference(cl, min_col=1, min_row=2, max_row=1 + n_sub))
cl.add_chart(chart, "F2")
cl.cell(row=n_sub + 3, column=1, value="緑の文字は「実験一覧」シートへの参照。LB−CV が正なら、LB が CV より悪い").font = f()

wb.calculation.fullCalcOnLoad = True  # openpyxl は計算結果を保存しないので、開いたときに再計算させる
wb.save(OUT)
print("saved", OUT, "rows", len(rows), "submitted", n_sub)
